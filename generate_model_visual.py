"""
Presentation: ML Model — visual-first, minimal code, presentation-friendly.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ─────────────────────────────────────────────────────────────────
INK      = RGBColor(0x0F, 0x17, 0x2A)   # near-black
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE= RGBColor(0xF8, 0xF9, 0xFF)
PURPLE   = RGBColor(0x7C, 0x3A, 0xED)
PURPLE_D = RGBColor(0x4C, 0x1D, 0x95)
PURPLE_L = RGBColor(0xED, 0xE9, 0xFE)
BLUE     = RGBColor(0x25, 0x63, 0xEB)
BLUE_L   = RGBColor(0xDB, 0xEA, 0xFE)
TEAL     = RGBColor(0x0D, 0x94, 0x88)
TEAL_L   = RGBColor(0xCC, 0xFB, 0xF1)
GREEN    = RGBColor(0x05, 0x96, 0x69)
GREEN_L  = RGBColor(0xD1, 0xFA, 0xE5)
ORANGE   = RGBColor(0xEA, 0x58, 0x0C)
ORANGE_L = RGBColor(0xFF, 0xED, 0xD5)
RED      = RGBColor(0xDC, 0x26, 0x26)
RED_L    = RGBColor(0xFE, 0xE2, 0xE2)
YELLOW   = RGBColor(0xD9, 0x77, 0x06)
YELLOW_L = RGBColor(0xFE, 0xF3, 0xC7)
GRAY     = RGBColor(0x6B, 0x72, 0x80)
GRAY_L   = RGBColor(0xF3, 0xF4, 0xF6)
GRAY_D   = RGBColor(0x1F, 0x29, 0x37)

W = Inches(13.33)
H = Inches(7.5)


# ── Primitives ───────────────────────────────────────────────────────────────
def prs_new():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p

def sl(prs): return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, c):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = c

def box(slide, l, t, w, h, fill, line=None, line_w=1.5):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line: s.line.color.rgb = line; s.line.width = Pt(line_w)
    else: s.line.fill.background()
    return s

def T(slide, text, l, t, w, h, size=14, bold=False, color=INK,
      align=PP_ALIGN.LEFT, italic=False, wrap=True):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color
    r.font.name = "Segoe UI"
    return tb

def Tm(slide, lines_styles, l, t, w, h, default_size=13,
       default_color=INK, align=PP_ALIGN.LEFT):
    """Multi-paragraph textbox. lines_styles = [(text, size, bold, color), ...]"""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(lines_styles):
        text = item[0]
        size  = item[1] if len(item) > 1 else default_size
        bold  = item[2] if len(item) > 2 else False
        color = item[3] if len(item) > 3 else default_color
        space = item[4] if len(item) > 4 else 4
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = "Segoe UI"
    return tb

def pill(slide, label, l, t, w, h, bg_c, fg_c=WHITE, size=12, bold=True):
    box(slide, l, t, w, h, bg_c)
    T(slide, label, l, t, w, h, size=size, bold=bold, color=fg_c,
      align=PP_ALIGN.CENTER)

def arrow_r(slide, l, t, length, color=GRAY, head=True):
    """Horizontal right arrow."""
    box(slide, l, t + Inches(0.07), length - Inches(0.12), Inches(0.04), color)
    if head:
        T(slide, "▶", l + length - Inches(0.22), t, Inches(0.22), Inches(0.22),
          size=10, color=color, align=PP_ALIGN.CENTER)

def arrow_d(slide, l, t, length, color=GRAY):
    """Vertical down arrow."""
    box(slide, l + Inches(0.06), t, Inches(0.04), length - Inches(0.12), color)
    T(slide, "▼", l, t + length - Inches(0.2), Inches(0.16), Inches(0.2),
      size=9, color=color, align=PP_ALIGN.CENTER)

def section_title(slide, title, accent=PURPLE):
    box(slide, 0, 0, W, Inches(0.9), PURPLE_D)
    box(slide, 0, Inches(0.9), W, Inches(0.06), accent)
    T(slide, title, Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.72),
      size=28, bold=True, color=WHITE)

def number_circle(slide, n, l, t, r, bg_c=PURPLE, size=22):
    box(slide, l, t, r, r, bg_c)
    T(slide, str(n), l, t, r, r, size=size, bold=True,
      color=WHITE, align=PP_ALIGN.CENTER)

def tensor_shape(slide, label, shape_text, l, t, w, h,
                 label_c=PURPLE_D, bg_c=PURPLE_L):
    box(slide, l, t, w, h, bg_c, line=PURPLE, line_w=1)
    T(slide, label, l + Inches(0.08), t + Inches(0.06),
      w - Inches(0.16), Inches(0.36), size=12, bold=True, color=label_c)
    T(slide, shape_text, l + Inches(0.08), t + Inches(0.38),
      w - Inches(0.16), Inches(0.28), size=10, color=GRAY,
      italic=True)

def feature_grid(slide, l, t, items, cols, cell_w, cell_h, colors):
    for i, name in enumerate(items):
        c = i % cols
        r = i // cols
        x = l + c * (cell_w + Inches(0.06))
        y = t + r * (cell_h + Inches(0.05))
        clr = colors[i % len(colors)]
        box(slide, x, y, cell_w, cell_h, clr)
        T(slide, name, x, y, cell_w, cell_h,
          size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 1 — COVER
# ═══════════════════════════════════════════════════════════════════════════
def s01_cover(prs):
    s = sl(prs); bg(s, INK)

    # Background accent shapes
    box(s, W - Inches(3.5), 0, Inches(3.5), H, RGBColor(0x1E, 0x20, 0x3A))
    box(s, 0, H - Inches(0.08), W, Inches(0.08), PURPLE)

    # Title block
    box(s, Inches(0.5), Inches(1.4), Inches(8.5), Inches(0.07), PURPLE)
    T(s, "Bagaimana Model Memprediksi", Inches(0.5), Inches(1.55),
      Inches(8.5), Inches(0.8), size=36, bold=True, color=WHITE)
    T(s, "Kualitas Udara?",
      Inches(0.5), Inches(2.35), Inches(8.5), Inches(0.8),
      size=36, bold=True, color=PURPLE)
    T(s, "Penjelasan arsitektur BiLSTM + LightGBM\nsecara visual, langkah per langkah",
      Inches(0.5), Inches(3.25), Inches(8.5), Inches(0.7),
      size=15, color=RGBColor(0xA5, 0xB4, 0xFC))

    # Right — key numbers
    nums = [
        ("24 jam", "data input per prediksi"),
        ("6 jam", "prediksi ke depan"),
        ("14 parameter", "polutan & cuaca"),
        ("3 output", "nilai + batas bawah + atas"),
    ]
    for i, (big, small) in enumerate(nums):
        y = Inches(1.4 + i * 1.45)
        box(s, W - Inches(3.2), y, Inches(2.7), Inches(1.2), PURPLE_D)
        box(s, W - Inches(3.2), y, Inches(0.07), Inches(1.2), PURPLE)
        T(s, big, W - Inches(3.05), y + Inches(0.12),
          Inches(2.5), Inches(0.55), size=26, bold=True, color=WHITE)
        T(s, small, W - Inches(3.05), y + Inches(0.65),
          Inches(2.5), Inches(0.42), size=12, color=RGBColor(0xA5,0xB4,0xFC))


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 2 — BIG PICTURE
# ═══════════════════════════════════════════════════════════════════════════
def s02_big_picture(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "Gambaran Besar  —  Dari Data Sensor ke Prediksi")

    # Main flow pipeline (horizontal)
    steps = [
        (BLUE,    "①\nDATA\nSENSOR",   "24 jam\nhistoris"),
        (TEAL,    "②\nPRA-\nPROSES",   "Bersihkan\n& normalisasi"),
        (ORANGE,  "③\nSSA",            "Pisahkan\ntrend & osilasi"),
        (PURPLE,  "④\nBiLSTM",         "Pelajari\npola waktu"),
        (GREEN,   "⑤\nLightGBM",       "Koreksi\n& interval"),
        (RED,     "⑥\nHASIL",          "6 langkah\n+ confidence"),
    ]

    for i, (c, title, sub) in enumerate(steps):
        x = Inches(0.3 + i * 2.15)
        box(s, x, Inches(1.15), Inches(1.9), Inches(2.1), c)
        T(s, title, x, Inches(1.2), Inches(1.9), Inches(1.5),
          size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        box(s, x, Inches(3.25), Inches(1.9), Inches(0.7), GRAY_L,
            line=c, line_w=1.5)
        T(s, sub, x + Inches(0.05), Inches(3.27), Inches(1.8), Inches(0.66),
          size=10, color=GRAY_D, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            T(s, "→", x + Inches(1.9), Inches(1.9),
              Inches(0.25), Inches(0.4), size=18, bold=True,
              color=c, align=PP_ALIGN.CENTER)

    # Output detail
    box(s, Inches(0.3), Inches(4.25), W - Inches(0.6), Inches(0.06), PURPLE)
    T(s, "Yang dihasilkan untuk setiap sensor, setiap jam:", Inches(0.3),
      Inches(4.4), Inches(5), Inches(0.4), size=13, bold=True, color=PURPLE_D)

    outputs = [
        (GREEN,   "Nilai Prediksi",   "PM2.5, PM10, AQI... untuk jam ke +1, +2, +3, +4, +5, +6"),
        (BLUE_L,  "Batas Bawah (P10)","Nilai minimum yang wajar — 80% confidence interval"),
        (ORANGE_L,"Batas Atas  (P90)","Nilai maksimum yang wajar — 80% confidence interval"),
    ]
    for i, (c, label, desc) in enumerate(outputs):
        x = Inches(0.3 + i * 4.3)
        box(s, x, Inches(4.9), Inches(4.1), Inches(0.5), c)
        box(s, x, Inches(5.4), Inches(4.1), Inches(1.6), GRAY_L,
            line=c, line_w=1.5)
        T(s, label, x + Inches(0.1), Inches(4.93),
          Inches(3.9), Inches(0.44), size=13, bold=True, color=WHITE
          if c not in (BLUE_L, ORANGE_L) else INK)
        T(s, desc, x + Inches(0.1), Inches(5.44),
          Inches(3.9), Inches(1.5), size=11, color=GRAY_D)

    # Key question
    box(s, Inches(0.3), Inches(7.05), W - Inches(0.6), Inches(0.38), PURPLE_L)
    T(s, "❓  Pertanyaan utama: Dengan melihat kondisi 24 jam terakhir, apa yang akan terjadi "
         "pada kualitas udara dalam 6 jam ke depan?",
      Inches(0.45), Inches(7.07), W - Inches(0.9), Inches(0.34),
      size=12, bold=False, color=PURPLE_D)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 3 — DATA INPUT
# ═══════════════════════════════════════════════════════════════════════════
def s03_input(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "① Data Input  —  Apa yang Diberikan ke Model?", accent=BLUE)

    # Left: sensor illustration
    box(s, Inches(0.3), Inches(1.1), Inches(4.5), Inches(6.2), BLUE_L,
        line=BLUE, line_w=1.5)
    T(s, "📡  Data Sensor", Inches(0.45), Inches(1.2),
      Inches(4.2), Inches(0.45), size=16, bold=True, color=BLUE)
    T(s, "Tabel t_loggers — dibaca setiap jam",
      Inches(0.45), Inches(1.65), Inches(4.2), Inches(0.3),
      size=11, color=GRAY)

    params = [
        "PM2.5", "PM2.5 koreksi", "PM10", "PM10 koreksi",
        "TSP", "TSP koreksi", "Noise (dBA)", "Temp (°C)",
        "Tekanan (mmHg)", "Kelembaban (%)", "AQI PM2.5",
        "AQI PM10", "AQI TSP", "AQI Index",
    ]
    clrs = [BLUE, BLUE, TEAL, TEAL, GREEN, GREEN,
            ORANGE, ORANGE, RED, RED, PURPLE, PURPLE, PURPLE, PURPLE]
    for i, (name, c) in enumerate(zip(params, clrs)):
        col = i % 2; row = i // 2
        x = Inches(0.4 + col * 2.15)
        y = Inches(2.05 + row * 0.6)
        box(s, x, y, Inches(1.95), Inches(0.44), c)
        T(s, f"  {name}", x, y, Inches(1.95), Inches(0.44),
          size=10, bold=True, color=WHITE)

    # Right: timeline illustration
    box(s, Inches(5.1), Inches(1.1), Inches(7.93), Inches(6.2), GRAY_L)
    T(s, "⏱  Jendela Waktu Input", Inches(5.25), Inches(1.2),
      Inches(7.6), Inches(0.45), size=16, bold=True, color=GRAY_D)

    # Timeline bar
    box(s, Inches(5.3), Inches(1.9), Inches(7.5), Inches(0.5), INK)
    T(s, "24 jam yang lalu", Inches(5.3), Inches(1.93),
      Inches(3), Inches(0.44), size=11, bold=True, color=WHITE)
    T(s, "SEKARANG →", Inches(10.3), Inches(1.93),
      Inches(2.3), Inches(0.44), size=11, bold=True,
      color=YELLOW, align=PP_ALIGN.RIGHT)

    # Hour blocks
    for i in range(24):
        x = Inches(5.3 + i * 0.312)
        c = BLUE if i % 2 == 0 else RGBColor(0x1D,0x4E,0xD8)
        box(s, x, Inches(2.5), Inches(0.3), Inches(0.6), c)
        if i % 6 == 0:
            T(s, f"-{24-i}h", x - Inches(0.1), Inches(3.15),
              Inches(0.5), Inches(0.25), size=8, color=GRAY,
              align=PP_ALIGN.CENTER)

    T(s, "Setiap kotak = 1 jam data sensor (rata-rata)", Inches(5.3),
      Inches(3.5), Inches(7.5), Inches(0.3), size=11, color=GRAY, italic=True)

    # Matrix shape
    box(s, Inches(5.1), Inches(3.95), Inches(7.93), Inches(0.06), BLUE)
    T(s, "Bentuk data setelah preprocessing:",
      Inches(5.25), Inches(4.1), Inches(7.6), Inches(0.35),
      size=13, bold=True, color=GRAY_D)

    box(s, Inches(5.25), Inches(4.55), Inches(3.5), Inches(1.5), WHITE,
        line=BLUE, line_w=2)
    T(s, "24", Inches(5.25), Inches(4.65), Inches(3.5), Inches(0.45),
      size=30, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    T(s, "langkah waktu (jam)", Inches(5.25), Inches(5.05),
      Inches(3.5), Inches(0.35), size=11, color=GRAY, align=PP_ALIGN.CENTER)
    T(s, "×", Inches(8.85), Inches(4.9), Inches(0.45), Inches(0.45),
      size=22, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
    box(s, Inches(9.35), Inches(4.55), Inches(3.5), Inches(1.5), WHITE,
        line=TEAL, line_w=2)
    T(s, "18", Inches(9.35), Inches(4.65), Inches(3.5), Inches(0.45),
      size=30, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    T(s, "fitur (14 sensor + 4 waktu)", Inches(9.35), Inches(5.05),
      Inches(3.5), Inches(0.35), size=11, color=GRAY, align=PP_ALIGN.CENTER)

    T(s, "4 fitur waktu: sin(jam), cos(jam), sin(hari), cos(hari)  →  agar model tahu pola pagi/malam/Senin/Minggu",
      Inches(5.1), Inches(6.2), Inches(7.93), Inches(0.4),
      size=11, color=GRAY, italic=True)
    T(s, "Mengapa sin/cos? Supaya jam 23 dan jam 0 dianggap 'berdekatan', bukan jauh.",
      Inches(5.1), Inches(6.6), Inches(7.93), Inches(0.35),
      size=11, color=PURPLE_D, bold=True)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 4 — PREPROCESSING
# ═══════════════════════════════════════════════════════════════════════════
def s04_preprocess(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "② Pra-Proses  —  Membersihkan Data Sebelum Dipakai", accent=TEAL)

    steps = [
        (BLUE,   "1",  "Resample ke 1 jam",
         "Data sensor masuk acak (menit-an)\n→ dirata-rata per jam",
         "Menit: 10:01, 10:07, 10:45...\n→ Jam 10:00 = rata-rata semua"),
        (RED,    "2",  "Hapus Anomali",
         "Nilai yang mencurigakan\n(>3σ dari median 24 jam) diganti\ndengan interpolasi linear",
         "Contoh: PM2.5 = 9999\n(sensor error) → dihapus\n→ diisi dari nilai sekitar"),
        (GREEN,  "3",  "Isi Gap & Ratakan",
         "Data kosong diisi interpolasi,\nlalu nilai ekstrem dipotong\ndi batas IQR ×1.5",
         "Potong outlier agar distribusi\ntidak didominasi nilai\nextrem 1-2 kejadian"),
        (ORANGE, "4",  "Normalisasi [0, 1]",
         "Setiap nilai sensor\ndiskala ke rentang 0-1\nper sensor (bukan global)",
         "PM2.5 sensor A: 0–500\nPM2.5 sensor B: 0–80\n→ skala berbeda per UID"),
    ]

    for i, (c, num, title, desc, example) in enumerate(steps):
        col = i % 2; row = i // 2
        x = Inches(0.3 + col * 6.5)
        y = Inches(1.1 + row * 2.95)

        # Main box
        box(s, x, y, Inches(0.55), Inches(2.65), c)
        T(s, num, x, y + Inches(0.85), Inches(0.55), Inches(0.55),
          size=24, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        box(s, x + Inches(0.55), y, Inches(5.95), Inches(0.52), c)
        T(s, title, x + Inches(0.68), y + Inches(0.07),
          Inches(5.7), Inches(0.4), size=15, bold=True, color=WHITE)

        box(s, x + Inches(0.55), y + Inches(0.52),
            Inches(2.95), Inches(2.13), WHITE, line=c, line_w=1.5)
        T(s, desc, x + Inches(0.65), y + Inches(0.6),
          Inches(2.75), Inches(2.0), size=12, color=GRAY_D)

        box(s, x + Inches(3.5), y + Inches(0.52),
            Inches(3.0), Inches(2.13), GRAY_L, line=c, line_w=1.5)
        T(s, "💡 Contoh:", x + Inches(3.62), y + Inches(0.6),
          Inches(2.8), Inches(0.28), size=10, bold=True, color=c)
        T(s, example, x + Inches(3.62), y + Inches(0.88),
          Inches(2.8), Inches(1.7), size=10, color=GRAY_D, italic=True)

    # Bottom note
    box(s, Inches(0.3), Inches(7.05), W - Inches(0.6), Inches(0.38), TEAL_L)
    T(s, "✅  Setelah pra-proses: data bersih, tanpa NaN, tanpa outlier, ternormalisasi [0,1] — siap masuk model",
      Inches(0.45), Inches(7.08), W - Inches(0.9), Inches(0.3),
      size=12, bold=True, color=TEAL)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 5 — SSA
# ═══════════════════════════════════════════════════════════════════════════
def s05_ssa(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "③ SSA — Memisahkan Tren dari Fluktuasi", accent=TEAL)

    # Analogy
    box(s, Inches(0.3), Inches(1.05), W - Inches(0.6), Inches(0.72), TEAL_L)
    T(s, "💡  Analogi: Bayangkan sinyal sensor seperti musik. SSA memisahkannya menjadi "
         "melodi utama (tren) dan ritme/getaran (osilasi) — model belajar keduanya secara terpisah.",
      Inches(0.45), Inches(1.08), W - Inches(0.9), Inches(0.66),
      size=12, color=TEAL)

    # Signal decomposition visual
    signals = [
        (BLUE,   "SINYAL ASLI", "PM2.5 mentah\nselama 24 jam",
         [30,32,35,38,40,38,36,34,32,35,38,42,45,43,40,38,36,34,32,30,29,31,33,35]),
        (ORANGE, "TREN", "Pola naik-turun jangka panjang\n(hasil SVD komponen ke-1)",
         [30,31,32,34,36,37,37,36,35,35,36,37,39,40,40,39,37,35,33,31,30,30,31,32]),
        (PURPLE, "OSILASI", "Fluktuasi & siklus harian\n(sinyal asli − tren)",
         [0,1,3,4,4,1,-1,-2,-3,0,2,5,6,3,0,-1,-1,-1,-1,-1,-1,1,2,3]),
    ]

    bar_w = Inches(0.22)
    bar_gap = Inches(0.02)
    for col_i, (c, title, desc, vals) in enumerate(signals):
        base_x = Inches(0.3 + col_i * 4.3)
        chart_h = Inches(1.5)
        chart_base_y = Inches(4.7)
        chart_top_y = Inches(3.2)

        box(s, base_x, Inches(1.85), Inches(4.0), Inches(0.42), c)
        T(s, title, base_x + Inches(0.1), Inches(1.88),
          Inches(3.8), Inches(0.36), size=13, bold=True, color=WHITE)
        T(s, desc, base_x + Inches(0.1), Inches(2.32),
          Inches(3.8), Inches(0.5), size=10, color=GRAY, italic=True)

        mn, mx = min(vals), max(vals)
        rng = mx - mn if mx != mn else 1
        for j, v in enumerate(vals):
            norm = (v - mn) / rng
            bh = Inches(0.15) + norm * chart_h
            bx = base_x + Inches(0.05) + j * (bar_w + bar_gap)
            by = chart_base_y - bh + Inches(0.05)
            box(s, bx, by, bar_w, bh, c)

        # Axis
        box(s, base_x + Inches(0.05), chart_base_y + Inches(0.05),
            Inches(3.85), Inches(0.025), GRAY)
        T(s, "jam ke-1", base_x + Inches(0.05), chart_base_y + Inches(0.1),
          Inches(1.2), Inches(0.2), size=8, color=GRAY)
        T(s, "jam ke-24", base_x + Inches(2.7), chart_base_y + Inches(0.1),
          Inches(1.2), Inches(0.2), size=8, color=GRAY)

    # = sign
    T(s, "=", Inches(4.35), Inches(3.7), Inches(0.5), Inches(0.5),
      size=30, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
    T(s, "+", Inches(8.65), Inches(3.7), Inches(0.5), Inches(0.5),
      size=30, bold=True, color=GRAY, align=PP_ALIGN.CENTER)

    # Result
    box(s, Inches(0.3), Inches(5.6), W - Inches(0.6), Inches(0.06), TEAL)
    box(s, Inches(0.3), Inches(5.7), W - Inches(0.6), Inches(1.65), GRAY_L)

    T(s, "Hasil untuk BiLSTM:", Inches(0.45), Inches(5.75),
      Inches(4), Inches(0.35), size=13, bold=True, color=GRAY_D)

    items = [
        (ORANGE, "14 fitur sensor asli", "→", "14 kolom TREN"),
        (PURPLE, "14 fitur sensor asli", "→", "14 kolom OSILASI"),
        (BLUE,   "4 fitur waktu (sin/cos)", "→", "tetap 4 kolom"),
    ]
    for i, (c, frm, arr, to) in enumerate(items):
        x = Inches(0.35 + i * 4.3)
        box(s, x, Inches(6.15), Inches(1.65), Inches(0.38), c)
        T(s, frm, x, Inches(6.17), Inches(1.65), Inches(0.34),
          size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        T(s, arr + " " + to, x + Inches(1.7), Inches(6.17),
          Inches(2.5), Inches(0.34), size=10, color=GRAY_D, bold=(i == 2))

    T(s, "Total: 14×2 + 4 = 32 fitur masuk ke BiLSTM  (dari semula 18)",
      Inches(0.45), Inches(6.6), W - Inches(0.9), Inches(0.35),
      size=12, bold=True, color=TEAL)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 6 — BiLSTM
# ═══════════════════════════════════════════════════════════════════════════
def s06_bilstm(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "④ BiLSTM  —  Membaca Pola Urutan Waktu", accent=PURPLE)

    # Analogy
    box(s, Inches(0.3), Inches(1.05), W - Inches(0.6), Inches(0.65), PURPLE_L)
    T(s, "💡  Analogi: BiLSTM seperti membaca buku dua kali — sekali dari halaman 1→24, "
         "sekali dari 24→1 — lalu menggabungkan pemahaman dari kedua arah.",
      Inches(0.45), Inches(1.08), W - Inches(0.9), Inches(0.58),
      size=12, color=PURPLE_D)

    # Left: layer diagram
    layers = [
        (BLUE,   "INPUT",                "32 fitur × 24 jam",           "(1, 24, 32)"),
        (PURPLE, "BiLSTM layer 1",       "Baca maju & mundur\n128 unit/arah", "(1, 24, 256)"),
        (PURPLE, "BiLSTM layer 2",       "Baca maju & mundur\n64 unit/arah",  "(1, 24, 128)"),
        (RGBColor(0x6D,0x28,0xD9),
                 "BiLSTM layer 3",       "Meringkas 24 jam\nmenjadi 1 vektor",  "(1, 256)"),
        (GREEN,  "Dense + Reshape",      "Hitung 6 prediksi\nuntuk 14 fitur",   "(1, 6, 14)"),
        (ORANGE, "OUTPUT",               "6 langkah × 14 fitur\n(masih ternormalisasi)", "(1, 6, 14)"),
    ]

    for i, (c, name, desc, shp) in enumerate(layers):
        y = Inches(1.85 + i * 0.88)
        box(s, Inches(0.3), y, Inches(3.0), Inches(0.72), c)
        T(s, name, Inches(0.38), y + Inches(0.06),
          Inches(2.84), Inches(0.35), size=12, bold=True, color=WHITE)
        T(s, desc, Inches(0.38), y + Inches(0.38),
          Inches(2.84), Inches(0.3), size=9, color=WHITE)
        box(s, Inches(3.3), y, Inches(1.8), Inches(0.72), PURPLE_L,
            line=c, line_w=1)
        T(s, shp, Inches(3.32), y + Inches(0.12),
          Inches(1.76), Inches(0.48), size=10, bold=True,
          color=PURPLE_D, align=PP_ALIGN.CENTER)
        if i < len(layers) - 1:
            T(s, "↓", Inches(1.65), y + Inches(0.72),
              Inches(0.7), Inches(0.2), size=11, bold=True,
              color=c, align=PP_ALIGN.CENTER)

    # Right: Bidirectional explanation
    box(s, Inches(5.5), Inches(1.85), Inches(7.53), Inches(2.8),
        WHITE, line=PURPLE, line_w=1.5)
    T(s, "Mengapa Bidirectional?", Inches(5.65), Inches(1.92),
      Inches(7.2), Inches(0.4), size=14, bold=True, color=PURPLE_D)

    dirs = [
        (BLUE,   "→ Maju (jam 1→24)",
         "Melihat bagaimana polutan berkembang dari awal ke akhir"),
        (PURPLE, "← Mundur (jam 24→1)",
         "Melihat konteks dari akhir jendela ke awal — menangkap pola yang baru terlihat di akhir"),
        (GREEN,  "⊕ Gabungan",
         "Kedua arah digabung → representasi lebih kaya dan akurat per jam"),
    ]
    for i, (c, title, desc) in enumerate(dirs):
        y = Inches(2.42 + i * 0.72)
        box(s, Inches(5.65), y, Inches(0.06), Inches(0.55), c)
        T(s, title, Inches(5.78), y + Inches(0.02),
          Inches(7.1), Inches(0.28), size=11, bold=True, color=c)
        T(s, desc, Inches(5.78), y + Inches(0.3),
          Inches(7.1), Inches(0.28), size=10, color=GRAY)

    # Training
    box(s, Inches(5.5), Inches(4.85), Inches(7.53), Inches(0.42), GRAY_D)
    T(s, "Cara Training", Inches(5.65), Inches(4.88),
      Inches(7.2), Inches(0.35), size=13, bold=True, color=WHITE)

    train_items = [
        ("Loss function", "MSE — meminimalkan rata-rata kuadrat error"),
        ("Optimizer",     "Adam — adaptif, tidak perlu atur learning rate manual"),
        ("Early stopping","Berhenti jika tidak ada perbaikan selama 20 epoch berturut-turut"),
        ("Validation",    "10% data dipakai untuk memantau overfitting"),
        ("Max epoch",     "100 — tapi rata-rata berhenti jauh lebih awal"),
    ]
    for i, (k, v) in enumerate(train_items):
        y = Inches(5.35 + i * 0.42)
        box(s, Inches(5.5), y, Inches(2.0), Inches(0.38), GRAY_L)
        T(s, k, Inches(5.58), y + Inches(0.06),
          Inches(1.84), Inches(0.28), size=10, bold=True, color=GRAY_D)
        T(s, v, Inches(7.55), y + Inches(0.06),
          Inches(5.4), Inches(0.28), size=10, color=GRAY)

    box(s, Inches(5.5), Inches(7.45), Inches(7.53), Inches(0.38), GRAY_L)
    T(s, "Model disimpan di: saved_models/{uid}/bilstm.keras — "
         "dimuat ke memori sekali lalu di-cache",
      Inches(5.62), Inches(7.47), Inches(7.3), Inches(0.34),
      size=10, color=GRAY, italic=True)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 7 — LightGBM
# ═══════════════════════════════════════════════════════════════════════════
def s07_lgbm(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "⑤ LightGBM  —  Koreksi + Interval Kepercayaan", accent=GREEN)

    # Why
    box(s, Inches(0.3), Inches(1.05), W - Inches(0.6), Inches(0.72), GREEN_L)
    T(s, "💡  BiLSTM sudah bagus untuk pola urutan, tapi tidak tahu konteks kalender. "
         "LightGBM memperbaiki prediksi dengan menambahkan: jam berapa sekarang? hari apa? bulan apa?",
      Inches(0.45), Inches(1.08), W - Inches(0.9), Inches(0.66),
      size=12, color=GREEN)

    # Flow
    inputs = [
        (PURPLE, "14 nilai BiLSTM\nper langkah waktu",    "hasil prediksi awal\n(masih ternormalisasi)"),
        (ORANGE, "Jam target\n(0 – 23)",                  "jam prediksi, bukan\njam sekarang"),
        (BLUE,   "Hari dalam seminggu\n(0 = Senin)",      "pola weekday vs\nweekend berbeda"),
        (TEAL,   "Bulan\n(1 – 12)",                       "pola musiman\nberbeda antar bulan"),
    ]

    T(s, "Input per langkah waktu (6 langkah = 6 baris):",
      Inches(0.3), Inches(1.88), Inches(6.0), Inches(0.3),
      size=12, bold=True, color=GRAY_D)

    for i, (c, title, sub) in enumerate(inputs):
        x = Inches(0.3 + i * 3.1)
        box(s, x, Inches(2.28), Inches(2.85), Inches(0.42), c)
        T(s, title, x + Inches(0.1), Inches(2.3),
          Inches(2.65), Inches(0.38), size=11, bold=True, color=WHITE)
        box(s, x, Inches(2.7), Inches(2.85), Inches(0.58), GRAY_L,
            line=c, line_w=1)
        T(s, sub, x + Inches(0.1), Inches(2.73),
          Inches(2.65), Inches(0.52), size=10, color=GRAY)

    # Arrow
    T(s, "↓ digabung menjadi 17 fitur per langkah × 6 langkah = (6, 17)",
      Inches(0.3), Inches(3.38), Inches(12), Inches(0.35),
      size=12, bold=True, color=GREEN)

    # 3 models
    box(s, Inches(0.3), Inches(3.82), W - Inches(0.6), Inches(0.42), GRAY_D)
    T(s, "3 Model LightGBM yang Dilatih Secara Terpisah:",
      Inches(0.45), Inches(3.85), Inches(12), Inches(0.35),
      size=13, bold=True, color=WHITE)

    models = [
        (GREEN,   "Model TITIK\n(lgbm.pkl)",
         "Prediksi nilai paling mungkin",
         "MSE loss — minimalkan\nrata-rata kuadrat error",
         "PM2.5 = 42.3 µg/m³"),
        (BLUE,    "Model BAWAH\n(lgbm_lower.pkl)",
         "Batas bawah interval 80%",
         "Quantile loss α=0.10 —\n10% data ada di bawah ini",
         "PM2.5 ≥ 38.1 µg/m³"),
        (ORANGE,  "Model ATAS\n(lgbm_upper.pkl)",
         "Batas atas interval 80%",
         "Quantile loss α=0.90 —\n10% data ada di atas ini",
         "PM2.5 ≤ 48.7 µg/m³"),
    ]

    for i, (c, name, purpose, method, example) in enumerate(models):
        x = Inches(0.3 + i * 4.3)
        y = Inches(4.35)
        box(s, x, y, Inches(4.0), Inches(0.55), c)
        T(s, name, x + Inches(0.1), y + Inches(0.06),
          Inches(3.8), Inches(0.43), size=13, bold=True, color=WHITE)
        box(s, x, y + Inches(0.55), Inches(4.0), Inches(2.65),
            WHITE, line=c, line_w=1.5)
        Tm(s, [
            (purpose, 12, True, GRAY_D, 6),
            ("", 4, False, WHITE, 2),
            ("Cara kerja:", 10, True, c, 2),
            (method, 10, False, GRAY, 4),
            ("Contoh output:", 10, True, c, 2),
            (example, 11, True, GRAY_D, 2),
        ], x + Inches(0.12), y + Inches(0.62),
           Inches(3.76), Inches(2.5))

    # Confidence interval illustration
    box(s, Inches(0.3), Inches(7.1), W - Inches(0.6), Inches(0.38), GREEN_L)
    T(s, "Interval 80% = dari batas bawah (P10) ke batas atas (P90). "
         "Artinya: model 80% yakin nilai aktual akan berada di antara dua garis tersebut.",
      Inches(0.45), Inches(7.13), W - Inches(0.9), Inches(0.34),
      size=12, color=GREEN)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 8 — OUTPUT
# ═══════════════════════════════════════════════════════════════════════════
def s08_output(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "⑥ Output  —  Apa yang Dihasilkan Model?", accent=RED)

    # Grid: 6 steps × columns
    T(s, "Untuk setiap sensor, setiap jam menghasilkan:",
      Inches(0.3), Inches(1.12), Inches(8), Inches(0.38),
      size=14, bold=True, color=GRAY_D)

    steps_label = ["+1 jam", "+2 jam", "+3 jam", "+4 jam", "+5 jam", "+6 jam"]
    col_colors  = [RED, ORANGE, YELLOW, GREEN, TEAL, BLUE]

    # Header row
    for i, (label, c) in enumerate(zip(steps_label, col_colors)):
        x = Inches(2.65 + i * 1.75)
        box(s, x, Inches(1.6), Inches(1.65), Inches(0.5), c)
        T(s, label, x, Inches(1.62), Inches(1.65), Inches(0.46),
          size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Row labels
    row_labels = ["AQI Index", "PM2.5", "PM10", "Noise", "Temp", "Kelembaban"]
    row_colors = [PURPLE, BLUE, TEAL, ORANGE, RED, GREEN]
    for j, (rl, rc) in enumerate(zip(row_labels, row_colors)):
        y = Inches(2.18 + j * 0.72)
        box(s, Inches(0.3), y, Inches(2.3), Inches(0.55), rc)
        T(s, rl, Inches(0.35), y + Inches(0.08),
          Inches(2.2), Inches(0.4), size=11, bold=True, color=WHITE)

        for i, c in enumerate(col_colors):
            x = Inches(2.65 + i * 1.75)
            alpha = int(255 * (1 - i * 0.08))
            box(s, x, y, Inches(1.65), Inches(0.55), GRAY_L)
            T(s, "##.#", x, y + Inches(0.08), Inches(1.65), Inches(0.4),
              size=13, bold=True, color=c, align=PP_ALIGN.CENTER)

    # Right: 3 outputs per cell
    box(s, Inches(12.8), Inches(1.6), Inches(0.45), Inches(4.55),
        GRAY_L, line=GRAY, line_w=1)
    T(s, "×\n3\nou\ntput", Inches(12.8), Inches(2.5),
      Inches(0.45), Inches(1.5), size=10, bold=True, color=GRAY,
      align=PP_ALIGN.CENTER)

    # Legend
    T(s, "Tiap sel berisi:", Inches(0.3), Inches(6.5),
      Inches(3), Inches(0.35), size=13, bold=True, color=GRAY_D)
    legend = [
        (GREEN,  "Nilai prediksi (titik)"),
        (BLUE,   "Batas bawah  P10"),
        (ORANGE, "Batas atas   P90"),
    ]
    for i, (c, lbl) in enumerate(legend):
        x = Inches(0.3 + i * 4.2)
        box(s, x, Inches(6.92), Inches(0.32), Inches(0.32), c)
        T(s, lbl, x + Inches(0.4), Inches(6.94),
          Inches(3.7), Inches(0.28), size=12, color=GRAY_D)

    # Total count
    box(s, Inches(0.3), Inches(7.15), W - Inches(0.6), Inches(0.35), RED_L)
    T(s, "Total per prediksi:  6 langkah  ×  14 fitur  ×  3 output  =  252 nilai  "
         "→ disimpan ke database dan ditampilkan di dashboard",
      Inches(0.45), Inches(7.18), W - Inches(0.9), Inches(0.3),
      size=11, bold=True, color=RED)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 9 — TRAINING vs INFERENCE
# ═══════════════════════════════════════════════════════════════════════════
def s09_train_vs_infer(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "Training vs Inference  —  Dua Mode Operasi Model", accent=ORANGE)

    divider_x = Inches(6.65)
    box(s, divider_x, Inches(1.05), Inches(0.06), H - Inches(1.05), GRAY_L)

    # TRAINING
    box(s, Inches(0.3), Inches(1.05), Inches(6.2), Inches(0.48), ORANGE)
    T(s, "🏋  TRAINING  —  Belajar dari Sejarah 1 Tahun",
      Inches(0.42), Inches(1.08), Inches(5.96), Inches(0.42),
      size=14, bold=True, color=WHITE)

    train_steps = [
        (BLUE,   "1", "Ambil 8.760 jam data\n(1 tahun dari t_loggers)"),
        (TEAL,   "2", "Preprocessing:\nresample → anomaly → clean → normalize"),
        (ORANGE, "3", "Buat 8.730 sequence\n(sliding window 24 jam)"),
        (ORANGE, "4", "Split 90% train\n10% validasi"),
        (PURPLE, "5", "Train BiLSTM\n(hingga 100 epoch, early stop)"),
        (PURPLE, "6", "Jalankan BiLSTM di train set\n→ hasilkan input LightGBM"),
        (GREEN,  "7", "Train 3 LightGBM\n(point, lower, upper)"),
        (GRAY,   "8", "Hitung MAE di val set\n→ simpan sebagai baseline"),
    ]

    for i, (c, n, txt_) in enumerate(train_steps):
        col = i % 2; row = i // 2
        x = Inches(0.3 + col * 3.1)
        y = Inches(1.63 + row * 1.3)
        box(s, x, y, Inches(0.38), Inches(1.1), c)
        T(s, n, x, y + Inches(0.3), Inches(0.38), Inches(0.4),
          size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        box(s, x + Inches(0.38), y, Inches(2.6), Inches(1.1), WHITE,
            line=c, line_w=1.5)
        T(s, txt_, x + Inches(0.48), y + Inches(0.1),
          Inches(2.4), Inches(0.9), size=10, color=GRAY_D)
        if i < len(train_steps) - 1 and i % 2 == 1:
            T(s, "↓", Inches(3.4), y + Inches(1.1),
              Inches(0.4), Inches(0.22), size=11, color=ORANGE,
              align=PP_ALIGN.CENTER)

    # INFERENCE
    box(s, Inches(6.85), Inches(1.05), Inches(6.18), Inches(0.48), PURPLE)
    T(s, "⚡  INFERENCE  —  Prediksi Real-time Setiap Jam",
      Inches(6.97), Inches(1.08), Inches(5.94), Inches(0.42),
      size=14, bold=True, color=WHITE)

    infer_steps = [
        (BLUE,   "1", "Ambil 72 jam terakhir\n(3× N_INPUT_HOURS sebagai buffer)"),
        (TEAL,   "2", "Preprocessing:\nresample → anomaly → clean → normalize"),
        (ORANGE, "3", "Ambil 24 jam terakhir\nsebagai jendela input"),
        (PURPLE, "4", "Terapkan SSA:\n(1,24,18) → (1,24,32)"),
        (PURPLE, "5", "BiLSTM predict:\n(1,24,32) → (1,6,14)"),
        (GREEN,  "6", "LightGBM refine:\n(6,17) → (6,14) × 3"),
        (RED,    "7", "Denormalize:\nnilai [0,1] → µg/m³, °C, dsb"),
        (GRAY,   "8", "Simpan ke DB\ndan tampilkan di dashboard"),
    ]

    for i, (c, n, txt_) in enumerate(infer_steps):
        col = i % 2; row = i // 2
        x = Inches(6.85 + col * 3.05)
        y = Inches(1.63 + row * 1.3)
        box(s, x, y, Inches(0.38), Inches(1.1), c)
        T(s, n, x, y + Inches(0.3), Inches(0.38), Inches(0.4),
          size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        box(s, x + Inches(0.38), y, Inches(2.55), Inches(1.1), WHITE,
            line=c, line_w=1.5)
        T(s, txt_, x + Inches(0.48), y + Inches(0.1),
          Inches(2.35), Inches(0.9), size=10, color=GRAY_D)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 10 — DRIFT DETECTION
# ═══════════════════════════════════════════════════════════════════════════
def s10_drift(prs):
    s = sl(prs); bg(s, OFF_WHITE)
    section_title(s, "Bonus: Deteksi Drift  —  Model Tahu Kalau Dirinya Tidak Akurat", accent=RED)

    # Analogy
    box(s, Inches(0.3), Inches(1.05), W - Inches(0.6), Inches(0.65), RED_L)
    T(s, "💡  Analogi: Model seperti dokter yang terus memeriksa diagnosisnya. "
         "Setiap jam, sistem membandingkan prediksi dengan nilai aktual. "
         "Kalau kesalahan melonjak, model segera retrain otomatis.",
      Inches(0.45), Inches(1.08), W - Inches(0.9), Inches(0.58),
      size=12, color=RED)

    # Flow
    flow = [
        (BLUE,   "Prediksi\ndibuat",         "+1 jam"),
        (TEAL,   "Waktu target\ntiba",        "+1 jam kemudian"),
        (GREEN,  "Ambil nilai\naktual sensor","dari t_loggers ±30 min"),
        (ORANGE, "Hitung\nMAE",              "error vs prediksi"),
        (RED,    "Bandingkan\nbaseline",      "MAE sekarang\n÷ MAE saat training"),
        (PURPLE, "Jika > 1.5×\nAuto-retrain", "background thread"),
    ]

    for i, (c, title, sub) in enumerate(flow):
        x = Inches(0.3 + i * 2.15)
        box(s, x, Inches(1.85), Inches(1.9), Inches(1.6), c)
        T(s, title, x + Inches(0.1), Inches(1.93),
          Inches(1.7), Inches(0.9), size=13, bold=True, color=WHITE,
          align=PP_ALIGN.CENTER)
        box(s, x, Inches(3.45), Inches(1.9), Inches(0.65), GRAY_L, line=c, line_w=1)
        T(s, sub, x + Inches(0.08), Inches(3.48),
          Inches(1.74), Inches(0.58), size=9, color=GRAY_D,
          align=PP_ALIGN.CENTER, italic=True)
        if i < len(flow) - 1:
            T(s, "→", x + Inches(1.9), Inches(2.45),
              Inches(0.25), Inches(0.38), size=16, bold=True,
              color=c, align=PP_ALIGN.CENTER)

    # Drift score formula
    box(s, Inches(0.3), Inches(4.35), W - Inches(0.6), Inches(0.06), RED)
    box(s, Inches(0.3), Inches(4.45), W - Inches(0.6), Inches(2.15), GRAY_L)
    T(s, "Formula Drift Score", Inches(0.5), Inches(4.55),
      Inches(5), Inches(0.4), size=14, bold=True, color=GRAY_D)

    box(s, Inches(0.5), Inches(5.0), Inches(5.5), Inches(0.65), WHITE,
        line=RED, line_w=2)
    T(s, "drift_score  =  MAE terkini  ÷  MAE baseline",
      Inches(0.6), Inches(5.06), Inches(5.3), Inches(0.52),
      size=14, bold=True, color=RED, align=PP_ALIGN.CENTER)

    # Thresholds
    thresholds = [
        (GREEN,  "< 1.0",   "Model lebih akurat dari training\nTidak ada tindakan"),
        (ORANGE, "1.0–1.5", "Akurasi sedikit turun\nStatus tetap 'ready', dipantau"),
        (RED,    "> 1.5",   "Drift terdeteksi!\nStatus → 'drifted', auto-retrain dimulai"),
    ]
    for i, (c, val, desc) in enumerate(thresholds):
        x = Inches(6.2 + i * 2.3)
        box(s, x, Inches(4.7), Inches(2.1), Inches(0.5), c)
        T(s, val, x, Inches(4.73), Inches(2.1), Inches(0.44),
          size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        box(s, x, Inches(5.2), Inches(2.1), Inches(1.42), WHITE, line=c, line_w=1.5)
        T(s, desc, x + Inches(0.1), Inches(5.26),
          Inches(1.9), Inches(1.3), size=10, color=GRAY_D)

    # Status badges
    box(s, Inches(0.3), Inches(6.65), W - Inches(0.6), Inches(0.42), GRAY_D)
    T(s, "Status Model yang Muncul di Dashboard:",
      Inches(0.45), Inches(6.68), Inches(4), Inches(0.35),
      size=12, bold=True, color=WHITE)

    badges = [("ready", GREEN), ("training", BLUE), ("drifted", ORANGE),
              ("error", RED), ("untrained", GRAY)]
    for i, (name, c) in enumerate(badges):
        x = Inches(4.8 + i * 1.72)
        box(s, x, Inches(6.68), Inches(1.55), Inches(0.35), c)
        T(s, name, x, Inches(6.7), Inches(1.55), Inches(0.31),
          size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 11 — SUMMARY / PENUTUP
# ═══════════════════════════════════════════════════════════════════════════
def s11_summary(prs):
    s = sl(prs); bg(s, INK)
    box(s, 0, 0, W, Inches(0.07), PURPLE)
    box(s, 0, H - Inches(0.07), W, Inches(0.07), PURPLE)

    T(s, "Ringkasan Arsitektur Model",
      Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.65),
      size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    summaries = [
        (BLUE,   "①  DATA",
         "24 jam × 14 sensor + 4 time features\n→ shape (1, 24, 18)"),
        (TEAL,   "②  PREPROCESS",
         "Resample → Anomaly removal → IQR clip\n→ Normalisasi [0,1] per UID"),
        (ORANGE, "③  SSA",
         "14 sensor → 28 komponen (tren+osilasi)\n→ shape (1, 24, 32) masuk BiLSTM"),
        (PURPLE, "④  BiLSTM",
         "3 lapis LSTM bidirectional\n→ shape (1, 6, 14) prediksi awal"),
        (GREEN,  "⑤  LightGBM",
         "3 model: point + P10 + P90\n→ Nilai akhir + confidence band 80%"),
        (RED,    "⑥  OUTPUT",
         "6 langkah × 14 fitur × 3 nilai\n= 252 prediksi per siklus jam"),
    ]

    for i, (c, title, desc) in enumerate(summaries):
        col = i % 3; row = i // 2
        x = Inches(0.4 + col * 4.3)
        y = Inches(1.15 + row * 2.85)
        box(s, x, y, Inches(4.0), Inches(0.5), c)
        box(s, x, y, Inches(0.07), Inches(2.55), c)
        T(s, title, x + Inches(0.2), y + Inches(0.06),
          Inches(3.7), Inches(0.4), size=15, bold=True, color=WHITE)
        box(s, x, y + Inches(0.5), Inches(4.0), Inches(2.05),
            RGBColor(0x1E, 0x1B, 0x4B))
        T(s, desc, x + Inches(0.2), y + Inches(0.62),
          Inches(3.7), Inches(1.85), size=12, color=RGBColor(0xC4, 0xB5, 0xFD))

    # Bottom
    box(s, Inches(0.3), Inches(7.12), W - Inches(0.6), Inches(0.35), PURPLE_D)
    T(s, "73 automated tests  ·  Satu set model per sensor UID  ·  "
         "Auto-retrain jika drift > 1.5×  ·  Confidence interval 80%",
      Inches(0.4), Inches(7.15), W - Inches(0.8), Inches(0.3),
      size=11, color=RGBColor(0xA5, 0xB4, 0xFC), align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════
def main():
    prs = prs_new()
    s01_cover(prs)
    s02_big_picture(prs)
    s03_input(prs)
    s04_preprocess(prs)
    s05_ssa(prs)
    s06_bilstm(prs)
    s07_lgbm(prs)
    s08_output(prs)
    s09_train_vs_infer(prs)
    s10_drift(prs)
    s11_summary(prs)

    out = r"C:\Users\nurch\OneDrive\Documents\project\air_quality_forecast\Air_Quality_Model_Visual.pptx"
    prs.save(out)
    print(f"Saved : {out}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
